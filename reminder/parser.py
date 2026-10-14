"""把上传的过保数据（Excel / CSV）解析成统一的记录列表。

优先按表头名称识别列；表头识别不到时，退回旧版“维保日历数据.xlsx”的固定列顺序：
代表处 | 最终客户 | 合同名称 | 过保时间 | 维保类型 | 合同号 | 金额 | 提醒天数 | 邮箱名单
"""
import csv
import datetime as dt
import io
import re

import openpyxl

# 字段 -> 可识别的表头别名（比较时忽略空格和大小写）
COLUMN_ALIASES = {
    "office": ["代表处", "办事处", "地区部"],
    "customer": ["最终客户", "客户", "客户名称", "最终用户"],
    "contract_name": ["合同名称", "项目名称"],
    "expiry": ["过保时间", "过保日期", "到期时间", "到期日期", "维保结束时间", "维保到期日"],
    "service_type": ["维保类型", "服务类型"],
    "contract_no": ["合同号", "合同编号", "合同编码"],
    "amount": ["金额", "合同金额", "维保金额"],
    "remind_days": ["提醒天数", "提前提醒天数"],
    "emails": ["邮箱名单", "邮箱", "收件人", "邮件名单"],
    "lang": ["邮件语言", "语言", "Language"],
    "customer_class": ["客户分类"],
    "signer": ["签约客户", "签约客户名称", "总代"],
    "tier2": ["二代", "二代名称"],
    "mode": ["通知模式"],
    "product_lines": ["产品明细（型号×报价项数量）", "产品明细(型号×报价项数量)", "产品明细", "产品型号"],
}

# 旧版 维保日历数据.xlsx 的列顺序
LEGACY_ORDER = ["office", "customer", "contract_name", "expiry", "service_type",
                "contract_no", "amount", "remind_days", "emails"]

REQUIRED = ["expiry", "contract_no"]

DATE_FORMATS = ["%Y/%m/%d", "%Y-%m-%d", "%Y.%m.%d", "%Y%m%d", "%Y年%m月%d日",
                "%Y/%m/%d %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M", "%Y-%m-%d %H:%M"]

EMAIL_SEPARATORS = re.compile(r"[;；,，\n]+")
PRODUCT_SEPARATORS = re.compile(r"[;；\n]+")


def parse_lang(value):
    """“English / 英文 / EN” 为英文，其余（含留空）为中文。"""
    text = str(value or "").strip().lower()
    return "en" if text.startswith(("en", "英")) else "zh"


def _norm(header):
    return re.sub(r"\s+", "", str(header or "")).lower()


def map_columns(header_row):
    """返回 {字段: 列下标}；表头不含必需列时返回 None。"""
    lookup = {}
    for field, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            lookup.setdefault(_norm(alias), field)
    mapping = {}
    for idx, cell in enumerate(header_row):
        field = lookup.get(_norm(cell))
        if field and field not in mapping:
            mapping[field] = idx
    if all(f in mapping for f in REQUIRED):
        return mapping
    return None


def parse_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, (int, float)):
        # Excel 序列号日期
        if 20000 < value < 80000:
            return dt.date(1899, 12, 30) + dt.timedelta(days=int(value))
        value = str(int(value))
    text = str(value).strip()
    for fmt in DATE_FORMATS:
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def parse_amount(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = re.sub(r"[,$￥¥\s]", "", str(value))
    try:
        return float(text)
    except ValueError:
        return None


def parse_int(value):
    if value is None or value == "":
        return None
    try:
        return int(float(str(value).strip()))
    except ValueError:
        return None


def split_emails(value):
    if value is None:
        return []
    return [e.strip() for e in EMAIL_SEPARATORS.split(str(value)) if e.strip()]


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def read_rows(filename, content):
    name = filename.lower()
    if name.endswith(".csv"):
        for encoding in ("utf-8-sig", "gbk"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError("无法识别 CSV 文件编码，请另存为 UTF-8 或 xlsx")
        return [row for row in csv.reader(io.StringIO(text))]
    if name.endswith((".xlsx", ".xlsm")):
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        # 待填写名单保存时可能停在“填写说明”页，优先读“提醒名单”工作表
        ws = wb["提醒名单"] if "提醒名单" in wb.sheetnames else wb.active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    raise ValueError("仅支持 .xlsx / .xlsm / .csv 文件（.xls 请先另存为 .xlsx）")


def parse_file(filename, content):
    """返回 (records, warnings)。"""
    return parse_rows(read_rows(filename, content))


def parse_rows(rows):
    """解析提醒清单（含表头的二维列表），返回 (records, warnings)。"""
    if not rows:
        return [], ["文件为空"]

    warnings = []
    mapping = map_columns(rows[0])
    if mapping is None:
        mapping = {f: i for i, f in enumerate(LEGACY_ORDER)}
        warnings.append("未识别到表头（需包含“过保时间”“合同号”列），已按旧版维保日历列顺序读取")

    def cell(row, field):
        idx = mapping.get(field)
        return row[idx] if idx is not None and idx < len(row) else None

    records = []
    for line_no, row in enumerate(rows[1:], start=2):
        if not any(v not in (None, "") for v in row):
            continue
        contract_no = _text(cell(row, "contract_no"))
        raw_expiry = cell(row, "expiry")
        expiry = parse_date(raw_expiry)
        if not contract_no:
            warnings.append(f"第 {line_no} 行缺少合同号，已跳过")
            continue
        if expiry is None:
            warnings.append(f"第 {line_no} 行（合同 {contract_no}）过保时间无效：{raw_expiry!r}，已跳过")
            continue
        records.append({
            "id": len(records),
            "row": line_no,
            "office": _text(cell(row, "office")),
            "customer": _text(cell(row, "customer")),
            "contract_name": _text(cell(row, "contract_name")),
            "expiry": expiry.isoformat(),
            "service_type": _text(cell(row, "service_type")),
            "contract_no": contract_no,
            "amount": parse_amount(cell(row, "amount")),
            "remind_days": parse_int(cell(row, "remind_days")),
            "emails": split_emails(cell(row, "emails")),
            "lang": parse_lang(cell(row, "lang")),
            "customer_class": _text(cell(row, "customer_class")),
            "signer": _text(cell(row, "signer")),
            "tier2": _text(cell(row, "tier2")),
            "mode": "channel" if "渠道" in _text(cell(row, "mode")) else "customer",
            "product_lines": [p.strip() for p in PRODUCT_SEPARATORS.split(_text(cell(row, "product_lines"))) if p.strip()],
        })
    return records, warnings
