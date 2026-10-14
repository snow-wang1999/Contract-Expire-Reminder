"""生成“待填写收件人”的 Excel：提醒名单 + 填写说明 + 产品明细 + 已排除 + 筛选条件。"""
import datetime as dt
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

LIST_COLUMNS = ["代表处", "最终客户", "合同名称", "过保时间", "维保类型", "合同号", "金额", "邮箱名单"]
EMAIL_COL = LIST_COLUMNS.index("邮箱名单") + 1

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="C7000B")
FILL_ME = PatternFill("solid", fgColor="FFF2CC")
EXAMPLE_FILL = PatternFill("solid", fgColor="F2F2F2")

PRODUCT_COLUMNS = [("合同号", "contract_no"), ("过保时间", "expiry"), ("服务类型", "service_type"),
                   ("最终客户", "customer"), ("续保产品", "product"), ("目录价", "amount"),
                   ("EOS时间", "eos"), ("资产状态", "asset_status"), ("产品线LV1", "lv1"), ("原始行号", "row")]

GUIDE = [
    "填写说明",
    "",
    "1. 只需要填写【提醒名单】工作表中黄色的“邮箱名单”一列，其他列请勿修改。",
    "2. 一行就是一封提醒邮件（合同号 + 过保时间 + 维保类型），同一合同的不同维保类型会分成多行，请分别填写。",
    "3. 多个收件人用英文分号 ; 分隔（中文分号 ；也能识别）。可以填邮箱地址，也可以填 Outlook 能识别的姓名/工号。",
    "4. 邮箱名单留空的行不会发送；不需要提醒的行也可以直接删除整行。",
    "5. 已带出邮箱的行是根据上次填写的同一客户自动带出的，请核对。",
    "6. 填好后保存为 .xlsx，回到网页第③步上传。",
    "",
    "示例（灰色为示例，不要复制到提醒名单里）：",
]

EXAMPLES = [
    ["南非代表处", "City of Tshwane", "PIN CL SA COT_Brilliantel_Services", "2027-03-31", "Hi-Care",
     "0Y071024000326", 46000, "zhangsan@huawei.com;lisi@huawei.com"],
    ["尼日利亚代表处", "Access Bank Plc", "Nigeria_Service Renewal", "2027-06-30", "Warranty",
     "0Y056625000932", 38000, "wangwu 00123456;zhaoliu@huawei.com"],
]


def _header(ws, columns):
    ws.append(columns)
    for cell in ws[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _date(value):
    try:
        return dt.date.fromisoformat(value)
    except (TypeError, ValueError):
        return value


def build(records, excluded, funnel):
    wb = Workbook()

    ws = wb.active
    ws.title = "提醒名单"
    _header(ws, LIST_COLUMNS + ["产品数", "产品明细"])
    for r in records:
        products = r.get("products", [])
        ws.append([r["office"], r["customer"], r["contract_name"], _date(r["expiry"]), r["service_type"],
                   r["contract_no"], round(r["amount"], 2), ";".join(r.get("emails", [])),
                   len(products), "；".join(p["product"] for p in products if p.get("product"))])
    for row in ws.iter_rows(min_row=2):
        row[3].number_format = "yyyy-mm-dd"
        row[6].number_format = "#,##0"
        row[EMAIL_COL - 1].fill = FILL_ME
    ws.cell(1, EMAIL_COL).fill = PatternFill("solid", fgColor="BF8F00")
    _widths(ws, [16, 28, 40, 12, 12, 18, 12, 50, 8, 60])

    guide = wb.create_sheet("填写说明")
    for line in GUIDE:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=14)
    start = guide.max_row + 1
    guide.append(LIST_COLUMNS)
    for cell in guide[start]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
    for ex in EXAMPLES:
        guide.append(ex)
        for cell in guide[guide.max_row]:
            cell.fill = EXAMPLE_FILL
    _widths(guide, [16, 28, 36, 12, 12, 18, 10, 44])

    detail = wb.create_sheet("产品明细")
    _header(detail, [c for c, _ in PRODUCT_COLUMNS])
    for r in records:
        for p in r.get("products", []):
            detail.append([p.get(k) for _, k in PRODUCT_COLUMNS])
    _widths(detail, [18, 12, 12, 28, 36, 12, 12, 10, 18, 9])

    out = wb.create_sheet("已排除")
    _header(out, ["排除原因"] + [c for c, _ in PRODUCT_COLUMNS])
    for p in excluded:
        out.append([p["reason"]] + [p.get(k) for _, k in PRODUCT_COLUMNS])
    _widths(out, [30, 18, 12, 12, 28, 36, 12, 12, 10, 18, 9])

    cond = wb.create_sheet("筛选条件")
    period = funnel["period"]
    yes = lambda b: "包含" if b else "不包含"
    for row in [
        ["提醒时间段", f"{period['start']} ~ {period['end']}"],
        ["目录价门槛", f"合同合计目录价 > ${period['min_amount']:,.0f}"],
        ["已 EOS 设备", yes(period["include_eos"])],
        ["资产状态为不再拓展", yes(period["include_no_expand"])],
        ["数据中心能源及关键供电", yes(period["include_dc_energy"])],
        [],
        ["原始数据行数", funnel["rows"]],
        ["时间段外", funnel["out_of_period"]],
        ["已 EOS", funnel["eos"]],
        ["不再拓展", funnel["no_expand"]],
        ["数据中心能源及关键供电", funnel["dc_energy"]],
        ["缺少合同号/日期", funnel["bad_row"]],
        ["合并后提醒条数", funnel["groups"]],
        ["目录价未超过门槛", funnel["below_min"]],
        ["最终提醒条数", funnel["records"]],
    ]:
        cond.append(row)
    _widths(cond, [24, 40])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
